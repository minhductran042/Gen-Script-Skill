"""Short-video script tool: read the team's Excel format, check structure, export back to Excel.

Usage:
    py -I kich_ban.py doc  <in.xlsx> <out.json>   # Excel (team format) -> JSON
    py -I kich_ban.py kiem <in.json>              # structural checks, exit 1 on any "Nặng" error
    py -I kich_ban.py xuat <in.json> <out.xlsx>   # JSON -> Excel (team format)

Format spec: shared/luat/format-kich-ban.md. Error codes: shared/luat/loi-kich-ban.md.
"""
import json
import re
import sys
import unicodedata

import openpyxl
from openpyxl.styles import Alignment, Font, PatternFill

sys.stdout.reconfigure(encoding="utf-8")

HEADERS = ["STT", "Tóm tắt - Title", "Địa điểm", "Mục đích", "Nhân vật",
           "Nội dung", "Biểu cảm", "Góc quay", "Text nổi bật", "Ghi chú"]
SCENE_KEYS = ["muc_dich", "nhan_vat", "noi_dung", "bieu_cam", "goc_quay", "text_noi_bat", "ghi_chu"]
REQUIRED_SCENE_KEYS = ["muc_dich", "nhan_vat", "noi_dung", "goc_quay"]
MAX_ACTORS = 6  # see shared/luat/dieu-kien-san-xuat.md
# Common Vietnamese words in their diacritic-stripped form (là, thôi, không, được...).
STRIPPED_WORDS = {"la", "thoi", "khong", "duoc", "nguoi", "nhung", "cua", "voi", "roi", "gi", "nay", "dang",
                  "nhieu", "minh", "toi", "phai", "biet", "muon", "lam", "trong", "quan", "y", "a", "di", "day"}

TITLE_RE = re.compile(r"^Kịch bản\s*(\d+)\s*:\s*(.*?)\s*\((?:Thời lượng dự kiến:\s*)?([^()]*)\)\s*$")
GIAY_RE = re.compile(r"^(.*?)\s*\((\d+\s*[–-]\s*\d+\s*s)\)\s*$")


def _s(v):
    return "" if v is None else str(v).strip()


# ---------------------------------------------------------------- doc (xlsx -> json)

def doc(path_in, path_out):
    ws = openpyxl.load_workbook(path_in).active
    scripts, cur, in_table = [], None, False
    for row in ws.iter_rows(min_row=1, max_col=10, values_only=True):
        cells = [_s(v) for v in row]
        a = cells[0]
        if not any(cells):
            continue
        m = TITLE_RE.match(a)
        if m:
            cur = {"so": int(m.group(1)), "ten": m.group(2), "thoi_luong": m.group(3).strip(),
                   "boi_canh": "", "dao_cu": "", "tom_tat": "", "title": "", "dia_diem": "", "canh": []}
            scripts.append(cur)
            in_table = False
            continue
        if cur is None:
            continue
        if a == "STT":
            in_table = True
            continue
        if not in_table:
            if a.startswith("Đạo cụ:"):
                cur["dao_cu"] = a[len("Đạo cụ:"):].strip()
            else:
                cur["boi_canh"] = re.sub(r"\s+", " ", a)
            continue
        # Scene row
        if cells[1]:
            parts = re.split(r"\n\s*Title\s*:\s*", cells[1], maxsplit=1)
            cur["tom_tat"] = parts[0].strip()
            cur["title"] = parts[1].strip() if len(parts) > 1 else ""
        if cells[2] and not cur["dia_diem"]:
            cur["dia_diem"] = cells[2]
        scene = dict(zip(SCENE_KEYS, [cells[3], cells[4], cells[5], cells[6], cells[7], cells[8], cells[9]]))
        gm = GIAY_RE.match(scene["muc_dich"])
        scene["giay"] = ""
        if gm:
            scene["muc_dich"], scene["giay"] = gm.group(1), gm.group(2).replace(" ", "")
        if cells[2] and cells[2] != cur["dia_diem"]:
            scene["dia_diem"] = cells[2]
        cur["canh"].append(scene)
    data = {"du_an": ws.title, "kich_ban": scripts}
    with open(path_out, "w", encoding="utf-8") as f:
        json.dump(data, f, ensure_ascii=False, indent=2)
    print(f"Đã đọc {len(scripts)} kịch bản, {sum(len(k['canh']) for k in scripts)} beat → {path_out}")


# ---------------------------------------------------------------- kiem (checks)

def _no_diacritics_line(line):
    """True when a 4+ word clause is typed without diacritics (likely a draft note in dialogue).

    Plain-ASCII clauses are legit for English terms ("Zoom out ra xa"), so it also needs 2+ words
    that are normally written with diacritics (la -> là, thoi -> thôi...).
    """
    for clause in re.split(r"[,.:;!?\"“”()]", line):
        words = re.findall(r"[A-Za-zÀ-ỹ]+", clause)
        if len(words) < 4 or not all(ord(ch) < 128 for w in words for ch in w):
            continue
        if sum(1 for w in words if w.lower() in STRIPPED_WORDS) >= 2:
            return True
    return False


def _actors(nhan_vat):
    text = re.sub(r"\(.*?\)", "", nhan_vat)
    if re.search(r"cả\s*\d+\s*người", text, re.I):
        return None  # group label, counted by other beats
    names = [n.strip() for n in re.split(r"[+,&]| và ", text) if n.strip()]
    return [re.sub(r"^voice\s+", "", n, flags=re.I) for n in names]


def kiem(path_in):
    with open(path_in, encoding="utf-8") as f:
        data = json.load(f)
    issues = []  # (level, code, where, message)

    def add(level, code, where, msg):
        issues.append((level, code, where, msg))

    long_lines = {}  # normalized line -> script number, to catch copy-paste across scripts
    for kb in data.get("kich_ban", []):
        so = kb.get("so", "?")
        tag = f"KB{so}"
        for key in ["ten", "thoi_luong", "dao_cu", "tom_tat", "title", "dia_diem"]:
            if not _s(kb.get(key)):
                add("Vừa", "NH-03", tag, f"thiếu trường '{key}'")
        canh = kb.get("canh", [])
        if not canh:
            add("Nặng", "NH-03", tag, "không có beat nào")
            continue
        first = canh[0]
        if "hook" not in _s(first.get("muc_dich")).lower() or not _s(first.get("giay")):
            add("Vừa", "HK-03", f"{tag} beat 1", "beat đầu phải ghi Mục đích có 'Hook' và mốc giây (giay)")
        if not _s(canh[-1].get("text_noi_bat")):
            add("Nhẹ", "KT-02", f"{tag} beat {len(canh)}", "beat cuối chưa có Text nổi bật")

        actors = set()
        for i, c in enumerate(canh, 1):
            where = f"{tag} beat {i}"
            for key in REQUIRED_SCENE_KEYS:
                if not _s(c.get(key)):
                    add("Vừa", "NH-03", where, f"trống cột '{key}'")
            names = _actors(_s(c.get("nhan_vat")))
            if names:
                actors.update(n.lower() for n in names)
            for line in _s(c.get("noi_dung")).splitlines():
                line = line.strip()
                if _no_diacritics_line(line):
                    add("Vừa", "NH-02", where, f"dòng không dấu, có thể là ghi chú nháp: {line[:60]}")
                norm = unicodedata.normalize("NFC", re.sub(r"\s+", " ", line)).lower()
                if len(norm) >= 40:
                    other = long_lines.setdefault(norm, so)
                    if other != so:
                        add("Nặng", "NH-01", where, f"đoạn trùng nguyên văn với kịch bản {other}: {line[:60]}…")
        if len(actors) > MAX_ACTORS:
            add("Nặng", "SX-01", tag, f"{len(actors)} nhân vật ({', '.join(sorted(actors))}) > {MAX_ACTORS} diễn viên")

    order = {"Nặng": 0, "Vừa": 1, "Nhẹ": 2}
    issues.sort(key=lambda x: order[x[0]])
    if not issues:
        print("kiem: không có lỗi cấu trúc.")
        return 0
    print("| Mức | Mã | Vị trí | Vấn đề |\n|---|---|---|---|")
    for lv, code, where, msg in issues:
        print(f"| {lv} | {code} | {where} | {msg} |")
    n_heavy = sum(1 for x in issues if x[0] == "Nặng")
    print(f"\nTổng: {len(issues)} lỗi, {n_heavy} Nặng.")
    return 1 if n_heavy else 0


# ---------------------------------------------------------------- xuat (json -> xlsx)

def xuat(path_in, path_out):
    with open(path_in, encoding="utf-8") as f:
        data = json.load(f)
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = _s(data.get("du_an")) or "Kịch bản"
    color = "FF1F1F1F"
    f_title = Font(name="Google Sans", size=8, bold=True, color=color)
    f_body = Font(name="Google Sans Text", size=8, italic=True, color=color)
    yellow = PatternFill("solid", fgColor="FFFFFF00")
    wrap = Alignment(wrap_text=True, vertical="top")
    ws.column_dimensions["F"].width = 32.25
    ws.column_dimensions["J"].width = 11.25

    r = 1
    for kb in data.get("kich_ban", []):
        ws.cell(r, 1, f"Kịch bản {kb['so']}: {kb['ten']} (Thời lượng dự kiến: {kb['thoi_luong']})").font = f_title
        r += 1
        if _s(kb.get("boi_canh")):
            ws.cell(r, 1, kb["boi_canh"]).font = f_body
            ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=26)
            r += 1
        ws.cell(r, 1, f"Đạo cụ: {_s(kb.get('dao_cu'))}").font = f_body
        r += 1
        for col, h in enumerate(HEADERS, 1):
            c = ws.cell(r, col, h)
            c.font, c.fill, c.alignment = f_body, yellow, wrap
        r += 1
        for i, sc in enumerate(kb.get("canh", [])):
            muc_dich = _s(sc.get("muc_dich"))
            if _s(sc.get("giay")):
                muc_dich = f"{muc_dich} ({sc['giay']})"
            values = [""] * 10
            if i == 0:
                values[0] = kb["so"]
                values[1] = f"{_s(kb.get('tom_tat'))}\n\n\n\nTitle: {_s(kb.get('title'))}"
                values[2] = _s(kb.get("dia_diem"))
            if _s(sc.get("dia_diem")):
                values[2] = sc["dia_diem"]
            values[3:] = [muc_dich, _s(sc.get("nhan_vat")), _s(sc.get("noi_dung")), _s(sc.get("bieu_cam")),
                          _s(sc.get("goc_quay")), _s(sc.get("text_noi_bat")), _s(sc.get("ghi_chu"))]
            for col, v in enumerate(values, 1):
                c = ws.cell(r, col, v if v != "" else None)
                c.font, c.alignment = f_body, wrap
            r += 1
    wb.save(path_out)
    print(f"Đã xuất {len(data.get('kich_ban', []))} kịch bản → {path_out}")


if __name__ == "__main__":
    if len(sys.argv) < 3 or sys.argv[1] not in {"doc", "kiem", "xuat"}:
        print(__doc__)
        sys.exit(2)
    cmd = sys.argv[1]
    if cmd == "doc":
        doc(sys.argv[2], sys.argv[3])
    elif cmd == "kiem":
        sys.exit(kiem(sys.argv[2]))
    else:
        xuat(sys.argv[2], sys.argv[3])
